#!/usr/bin/env python3
"""
Imports the Knowledge Content Register (owner's workbook, 23 Oct 2026) into data/knowledge/register.json.

  python3 scripts/import_knowledge_register.py data/knowledge/source/K1_Verification_Progress_2026-10-23.xlsx

Nothing here publishes anything. Every entry is imported with its verification state as the workbook records it and
`publication.approved = false`; the site shows an entry only when the publication gate in src/lib/knowledge/schema.ts
passes (verified source + exact locator + Greek and English text + editorial review + explicit approval). Draft texts
from the K1 sheets are kept as `draft` — internal, never rendered.
Where the workbook's sheets disagree, the stricter state wins (the K1 «Evidence Gate» sheet over the register).
"""
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SECTION = {"Λεξικό": "term", "Μάθε τα βασικά": "article", "Συχνές Ερωτήσεις": "faq"}
TOPIC = {
    "Φορολογία": "forologia", "Επιχειρηματικότητα": "epixeirimatikotita", "Οικονομικά & Χρηματοδότηση": "oikonomika",
    "Εργασία & Ασφάλιση": "ergasia", "Επιδοτήσεις & Προγράμματα": "epidotiseis", "Ακίνητα & Στέγαση": "akinita",
    "Κοινωνικές Παροχές": "koinonikes-paroxes", "Νομοθεσία & Δημόσιο": "nomothesia",
}
REGISTER_STATUS = {
    "Προς επαλήθευση": "unverified",
    "Προσχέδιο — όχι εγκεκριμένο": "draft",
    "Προσχέδιο με επίσημη πηγή — προς τελικό έλεγχο": "source_located",
}
GATE_STATUS = {
    "NOT VERIFIED": "draft",
    "OFFICIAL PRIMARY SOURCE LOCATED — editorial draft": "source_located",
    "Official source + draft checked": "source_located_draft_checked",
}
RANK = {"unverified": 0, "draft": 1, "source_located": 2, "source_located_draft_checked": 3}


def s(v):
    return None if v is None or str(v).strip() == "" else str(v).strip()


def main(path: str) -> None:
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = [r for r in wb["Content Register"].iter_rows(min_row=2, values_only=True) if r[0]]
    k1 = {r[1]: r for r in wb["K1 — Προτεραιότητες"].iter_rows(min_row=2, values_only=True) if r[1]}
    gate = {r[0]: r for r in wb["K1 — Evidence Gate"].iter_rows(min_row=2, values_only=True) if r[0]}
    researched = {r[0]: r for r in wb["K1 — Ερευνημένα προσχέδια"].iter_rows(min_row=2, values_only=True) if r[0]}
    entries = []
    disagreements = []
    for r in rows:
        (rid, section, title_el, topic, title_en, priority, status, authority, url, locator, checked, approved, note, related) = r[:14]
        en = s(title_en)
        if en and ("εκκρεμεί" in en.lower() or en.upper().startswith("EL / EN")):
            en = None
        state = REGISTER_STATUS[status]
        e = {
            "id": rid,
            "kind": SECTION[section],
            "topic": TOPIC[topic],
            "title": {"el": s(title_el), "en": en},
            "priority": priority,
            "body": None,
            "sources": [],
            "verification": {"status": state, "checked_on": s(checked), "candidate_authority": s(authority),
                             "starting_point": s(url), "locator_note": s(locator), "note": s(note)},
            "editorial_review": None,
            "publication": {"approved": False, "approved_by": None, "approved_on": None},
            "related": [x.strip() for x in str(related).split(",")] if related else [],
            "draft": None,
        }
        if str(approved).strip() not in ("Όχι", "None", ""):
            raise SystemExit(f"{rid}: the workbook marks it approved ({approved!r}) — refusing to import an approval without review")
        g = gate.get(rid)
        if g:
            gstate = GATE_STATUS[g[2]]
            if gstate != state:
                disagreements.append({"id": rid, "register": state, "evidence_gate": gstate})
            # the stricter of the two: a «source located» in one sheet is not enough if the other says less
            e["verification"]["status"] = min(state, gstate, key=lambda x: RANK[x]) if state != "unverified" else state
            e["verification"]["evidence_gate"] = g[2]
            if s(g[4]):
                e["verification"]["official_source_candidate"] = {"url": s(g[4]), "scope": s(g[5]), "amendment": s(g[8])}
            e["verification"]["remaining"] = s(g[9])
        k = k1.get(rid)
        rd = researched.get(rid)
        if k or rd:
            e["draft"] = {
                "el": s(rd[3]) if rd else s(k[8]),
                "en": s(rd[4]) if rd else s(k[9]),
                "status": (s(rd[8]) if rd else s(k[10])),
                "note": "Internal draft from the K1 workbook — not verified, not approved, never rendered.",
            }
            e["k1_order"] = int(k[0]) if k else None
        entries.append(e)
    ids = [e["id"] for e in entries]
    assert len(ids) == len(set(ids)), "duplicate ids"
    out = {
        "schema": "jackal-knowledge-register/1",
        "source_workbook": Path(path).name,
        "imported_on": "2026-10-23",
        "note": "Register of planned Knowledge entries. Nothing here is published; see the publication gate.",
        "sheet_disagreements": disagreements,
        "entries": entries,
    }
    dest = ROOT / "data/knowledge/register.json"
    dest.write_text(json.dumps(out, ensure_ascii=False, indent=1) + "\n")
    print(len(entries), "entries;", len(disagreements), "sheet disagreements")


if __name__ == "__main__":
    main(sys.argv[1])
